import io
import re
import json
import base64
from pathlib import Path
from typing import List, Dict
import datetime
import threading
import sys
import time

import google.generativeai as genai
import PyPDF2
import pandas as pd
import docx
from PIL import Image


API_KEY = "AIzaSyDRVj0Z_4xI3fy55Ux657tHnNGoC7JHl5g"
genai.configure(api_key=API_KEY)

def spinner(stop_event):
    """A simple console spinner to indicate processing."""
    spinner_chars = ['|', '/', '-', '\\']
    idx = 0
    while not stop_event.is_set():
        sys.stdout.write(f"\rGenerating document... {spinner_chars[idx % len(spinner_chars)]}")
        sys.stdout.flush()
        idx += 1
        time.sleep(0.1)
    sys.stdout.write("\rGeneration complete!             \n")

class MultiMediaProcessor:
    def __init__(self):
        """
        Initialize the MultiMediaProcessor with Gemini 2.0 Flash API.
        """
        self.model = genai.GenerativeModel(
            'gemini-2.0-flash',
            generation_config={
                'temperature': 0.3,  # Balanced creativity and factuality
                'top_p': 0.9,
                'max_output_tokens': 15000  # Increase if needed
            },
            safety_settings={
                'HARASSMENT': 'block_none',
                'HATE': 'block_none',
                'SEXUAL': 'block_none',
                'DANGEROUS': 'block_none'
            }
        )
        # These attributes will store the original source information
        self.original_source_content = ""
        self.original_source_names = []
        self.document_name = ""
        self.business_domain = None

    def _encode_image(self, image_path: str) -> str:
        """Encode an image to base64."""
        with open(image_path, "rb") as image_file:
            return base64.b64encode(image_file.read()).decode('utf-8')

    def extract_text(self, file_path: str) -> str:
        """Extract text from various file types."""
        path = Path(file_path)
        try:
            if path.suffix.lower() in ['.jpg', '.jpeg', '.png', '.gif', '.bmp']:
                return self._process_image(path)
            elif path.suffix.lower() == '.pdf':
                return self._extract_pdf_text(path)
            elif path.suffix.lower() == '.docx':
                return self._extract_docx_text(path)
            elif path.suffix.lower() in ['.xls', '.xlsx']:
                return self._extract_excel_text(path)
            elif path.suffix.lower() == '.txt':
                return self._extract_txt_text(path)
            else:
                print(f"Unsupported file type: {path.suffix}")
                return ""
        except Exception as e:
            print(f"Error extracting text from {file_path}: {e}")
            return ""

    def _process_image(self, path: Path) -> str:
        """Process and extract text from an image."""
        with Image.open(path) as img:
            max_size = (800, 800)
            img.thumbnail(max_size)
            buffer = io.BytesIO()
            img.save(buffer, format=img.format)
            buffer.seek(0)
        image_base64 = base64.b64encode(buffer.getvalue()).decode('utf-8')
        try:
            response = self.model.generate_content([
                "Thoroughly describe the contents of this image. "
                "Extract any text visible in the image. "
                "Provide a comprehensive analysis with any relevant insights.",
                {
                    'mime_type': 'image/jpeg',
                    'data': image_base64
                }
            ])
            return response.text
        except Exception as e:
            print(f"Error processing image {path}: {e}")
            return ""

    def _extract_pdf_text(self, path: Path) -> str:
        """Extract text from a PDF."""
        with path.open('rb') as file:
            reader = PyPDF2.PdfReader(file)
            return '\n'.join(page.extract_text() for page in reader.pages)

    def _extract_docx_text(self, path: Path) -> str:
        """Extract text from a DOCX file."""
        doc = docx.Document(path)
        return '\n'.join(paragraph.text for paragraph in doc.paragraphs)

    def _extract_excel_text(self, path: Path) -> str:
        """Extract text from an Excel file."""
        df = pd.read_excel(path)
        return df.to_string()

    def _extract_txt_text(self, path: Path) -> str:
        """Extract text from a plain text file."""
        return path.read_text(encoding='utf-8')

    def process_multiple_documents(self,
                                   document_type: str,
                                   source_files: List[str],
                                   business_domain: str = None) -> Dict[str, str]:
        """
        Combine the content of all source files and generate one document.
        """
        combined_content = ""
        source_names = []

        for file_path in source_files:
            try:
                extracted_text = self.extract_text(file_path)
                file_name = Path(file_path).name
                combined_content += f"\n--- Content from {file_name} ---\n{extracted_text}\n"
                source_names.append(file_name)
            except Exception as e:
                print(f"Error processing {file_path}: {e}")

        # Store original source content and related details for later regeneration
        self.original_source_content = combined_content
        self.original_source_names = source_names

        # Generate document name based on date and type
        current_date = datetime.datetime.now().strftime("%Y%m%d")
        self.document_name = f"{document_type}_{current_date}"
        self.business_domain = business_domain

        # Generate document content in sections
        generated_doc = self._generate_document_in_sections(
            document_type,
            combined_content,
            self.document_name,
            source_names,
            business_domain
        )

        return {self.document_name: generated_doc}

    def _generate_document_in_sections(self,
                                         document_type: str,
                                         source_content: str,
                                         document_name: str,
                                         source_names: List[str],
                                         business_domain: str = None) -> str:
        """
        Generate the document in sections to minimize token truncation.
        The generated document uses markdown formatting with bold headers.
        """
        if document_type.upper() == 'BRD':
            sections = [
                "Executive Summary",
                "Project Objectives",
                "Project Scope",
                "Business Requirements",
                "Key Stakeholders",
                "Project Constraints",
                "Cost-Benefit Analysis",
                "MoSCoW Prioritization Table",
                "References"
            ]
            prompt_template = (
                "Using the following source content and internal references, generate the section titled '{section}' "
                "for the Business Requirements Document named '{document_name}'. Ensure to include industry best practices, "
                "market trends, and at least some external references where applicable.\n\n"
                "Business Domain Context: {business_context}\n\n"
                "Source Content:\n{source_content}\n\n"
                "Internal Sources: {internal_sources}\n"
            )
        elif document_type.upper() == 'SRS':
            sections = [
                "Introduction",
                "General Description",
                "Functional Requirements",
                "Interface Requirements",
                "Performance Requirements",
                "Design Constraints",
                "Non-Functional Attributes",
                "MoSCoW Requirements Prioritization",
                "Preliminary Schedule and Budget",
                "References"
            ]
            prompt_template = (
                "Using the following source content and internal references, generate the section titled '{section}' "
                "for the Software Requirements Specification (SRS) named '{document_name}'. Ensure to include technical best practices, "
                "current technology insights, and at least some external references where applicable.\n\n"
                "Business Domain Context: {business_context}\n\n"
                "Source Content:\n{source_content}\n\n"
                "Internal Sources: {internal_sources}\n"
            )
        else:
            raise ValueError("Document type must be 'BRD' or 'SRS'")

        business_context = business_domain if business_domain else "General"
        internal_sources = ", ".join(source_names)

        sections_content = []
        for section in sections:
            prompt = prompt_template.format(
                section=section,
                document_name=document_name,
                business_context=business_context,
                source_content=source_content,
                internal_sources=internal_sources
            )
            try:
                response = self.model.generate_content(prompt)
                # Bolden the section header using markdown formatting
                section_text = f"## **{section}**\n\n{response.text}\n"
                sections_content.append(section_text)
            except Exception as e:
                print(f"Error generating section {section}: {e}")
                sections_content.append(f"## **{section}**\n\nError generating content for this section.\n")
        generated_doc = "\n".join(sections_content)
        return generated_doc

    def save_documents(self,
                      documents: Dict[str, str],
                      output_directory: str) -> None:
        """
        Save generated documents as PDF, text, and markdown files.
        """
        output_path = Path(output_directory)
        output_path.mkdir(parents=True, exist_ok=True)

        for filename, content in documents.items():
            # Save as text file for backup
            txt_output_filename = f"{filename}_with_references.txt"
            txt_file_path = output_path / txt_output_filename
            with txt_file_path.open('w', encoding='utf-8') as f:
                f.write(content)

            # Save as markdown file
            md_output_filename = f"{filename}_with_references.md"
            md_file_path = output_path / md_output_filename
            with md_file_path.open('w', encoding='utf-8') as f:
                f.write(content)

            # Save as PDF
            pdf_output_filename = f"{filename}_with_references.pdf"
            pdf_file_path = output_path / pdf_output_filename
            self._convert_to_pdf(content, str(pdf_file_path))

            print(f"Saved document as PDF: {pdf_file_path}")
            print(f"Backup text file saved: {txt_file_path}")
            print(f"Markdown file saved: {md_file_path}")

    def _convert_to_pdf(self, content: str, output_path: str) -> None:
        """
        Convert text content to PDF using FPDF with improved table handling.
        """
        try:
            from fpdf import FPDF

            pdf = FPDF()
            pdf.add_page()
            pdf.set_auto_page_break(auto=True, margin=15)
            pdf.set_font('Helvetica', '', 11)

            lines = content.split('\n')
            in_table = False
            table_buffer = []
            i = 0
            while i < len(lines):
                line = lines[i].strip()

                # Check for table indicators (e.g., MoSCoW table)
                if "MOSCOW" in line.upper() or (in_table and line):
                    in_table = True
                    if not table_buffer and "MOSCOW" in line.upper():
                        pdf.set_font('Helvetica', 'B', 14)
                        pdf.ln(10)
                        pdf.cell(0, 10, line, 0, 1)
                        pdf.set_font('Helvetica', '', 11)
                        pdf.ln(5)
                    elif "|" in line or "+-" in line or "-+-" in line or line.startswith('+'):
                        table_buffer.append(line)
                    else:
                        if table_buffer:
                            pdf.set_font('Courier', '', 9)
                            for table_line in table_buffer:
                                pdf.cell(0, 5, table_line, 0, 1)
                            pdf.set_font('Helvetica', '', 11)
                            pdf.ln(5)
                            table_buffer = []
                            in_table = False
                            continue
                        else:
                            in_table = False
                elif not line:
                    if in_table and table_buffer:
                        pdf.set_font('Courier', '', 9)
                        for table_line in table_buffer:
                            pdf.cell(0, 5, table_line, 0, 1)
                        pdf.set_font('Helvetica', '', 11)
                        pdf.ln(5)
                        table_buffer = []
                        in_table = False
                    else:
                        pdf.ln(5)
                elif line.startswith('#') or line.upper() == line or line.endswith(':'):
                    pdf.set_font('Helvetica', 'B', 14)
                    pdf.ln(10)
                    pdf.cell(0, 10, line, 0, 1)
                    pdf.set_font('Helvetica', '', 11)
                    pdf.ln(5)
                else:
                    pdf.multi_cell(0, 5, line)
                i += 1

            if table_buffer:
                pdf.set_font('Courier', '', 9)
                for table_line in table_buffer:
                    pdf.cell(0, 5, table_line, 0, 1)

            pdf.output(output_path)
        except Exception as e:
            print(f"Error creating PDF: {e}")
            print("Attempting simplified PDF creation...")
            try:
                from fpdf import FPDF
                pdf = FPDF()
                pdf.add_page()
                pdf.set_font('Courier', '', 10)
                pdf.set_auto_page_break(auto=True, margin=15)
                for line in content.split('\n'):
                    if line.strip():
                        pdf.multi_cell(0, 5, line)
                    else:
                        pdf.ln(5)
                pdf.output(output_path)
                print("Created simplified PDF as fallback")
            except Exception as e2:
                print(f"Fallback PDF creation failed: {e2}")
                print("Only text file will be available")

def load_document(file_path: str) -> str:
    """
    Load a document (text or markdown) from file.
    """
    path = Path(file_path)
    if path.exists():
        return path.read_text(encoding='utf-8')
    else:
        raise FileNotFoundError(f"{file_path} does not exist.")

def ai_assisted_edit(processor: MultiMediaProcessor, document_content: str, additional_context: str) -> str:
    """
    Use the AI model to refine the document content with additional context.
    """
    prompt = (
        "Below is the current document content:\n\n"
        f"{document_content}\n\n"
        "Additional context and instructions:\n"
        f"{additional_context}\n\n"
        "Please refine the document with the new context in mind. "
        "Provide an updated version that improves clarity and incorporates the additional details."
    )
    try:
        response = processor.model.generate_content(prompt)
        return response.text
    except Exception as e:
        print(f"Error during AI-assisted edit: {e}")
        return document_content  # Fallback to original content if an error occurs

def iterative_editing_loop(processor: MultiMediaProcessor, file_path: str, document_type: str):
    """
    Load a document and allow iterative AI-assisted and manual edits.
    Includes an option to regenerate the document from original sources.
    """
    current_content = load_document(file_path)
    print("\nCurrent Document Content:\n")
    print(current_content)
    
    while True:
        print("\nOptions:")
        print("1. Make manual edits and update the document file.")
        print("2. Provide additional context for AI-assisted revision.")
        print("3. Regenerate the document from original sources.")
        print("4. Exit iterative editing.")
        choice = input("Enter your choice (1/2/3/4): ").strip()
        
        if choice == '1':
            print("\nPlease open the document in your favorite editor, make changes, and save the file.")
            input("Press Enter after you have updated the document...")
            current_content = load_document(file_path)
            print("Document updated manually.")
            
        elif choice == '2':
            additional_context = input("Enter additional context/instructions for the AI: ").strip()
            revised_content = ai_assisted_edit(processor, current_content, additional_context)
            print("\nAI suggested the following revision:\n")
            print(revised_content)
            accept = input("\nAccept changes? (y/n): ").strip().lower()
            if accept == 'y':
                current_content = revised_content
                with open(file_path, 'w', encoding='utf-8') as f:
                    f.write(current_content)
                print("Document updated with AI revisions.")
            else:
                print("No changes were made.")
                
        elif choice == '3':
            print("Regenerating document from original sources...")
            new_doc = processor._generate_document_in_sections(
                document_type,
                processor.original_source_content,
                processor.document_name,
                processor.original_source_names,
                processor.business_domain
            )
            current_content = new_doc
            with open(file_path, 'w', encoding='utf-8') as f:
                f.write(current_content)
            print("Document regenerated.")
            
        elif choice == '4':
            print("Exiting iterative editing.")
            break
        else:
            print("Invalid choice. Please select 1, 2, 3, or 4.")

    return current_content

def main():
    document_type = input("Enter document type (BRD/SRS): ").strip().upper()
    while document_type not in ['BRD', 'SRS']:
        print("Invalid document type. Please enter BRD or SRS.")
        document_type = input("Enter document type (BRD/SRS): ").strip().upper()

    business_domain = input("Enter business domain (e.g., healthcare, finance, retail) or press Enter to skip: ").strip()
    if not business_domain:
        business_domain = None

    print("\nUsing local file system.")
    source_files = []
    print("\nEnter paths to source files (supports PDF, DOCX, XLSX, TXT, Images)")
    print("Tip: Use full path, e.g., '/content/file.pdf'")
    print("Press Enter twice to finish.")

    while True:
        file_path = input("File path: ").strip()
        if not file_path:
            if not source_files:
                print("No files entered. Exiting.")
                return
            break
        path = Path(file_path)
        if not path.exists():
            print(f"Error: File {file_path} does not exist.")
            continue
        source_files.append(str(path))

    output_directory = "./generated_documents_with_references"
    Path(output_directory).mkdir(parents=True, exist_ok=True)
    print(f"\nDocuments will be saved to: {output_directory}")

    try:
        # Check if fpdf is installed
        try:
            import fpdf
        except ImportError:
            print("Installing required PDF library (fpdf)...")
            import subprocess
            subprocess.check_call(["pip", "install", "fpdf"])
            print("FPDF installed successfully!")

        processor = MultiMediaProcessor()
        
        # Start spinner thread while generating document
        stop_spinner = threading.Event()
        spinner_thread = threading.Thread(target=spinner, args=(stop_spinner,))
        spinner_thread.start()
        
        generated_documents = processor.process_multiple_documents(
            document_type,
            source_files,
            business_domain
        )
        
        # Stop spinner
        stop_spinner.set()
        spinner_thread.join()

        processor.save_documents(generated_documents, output_directory)
        print(f"\nSuccessfully processed {len(generated_documents)} {document_type} document with references!")
        print(f"Internal sources: {len(source_files)} files")
        print("External references have been added to the document.")
        print("MoSCoW prioritization table included in the document.")
        print("Document saved as PDF, text backup, and markdown file.")

        # Offer iterative editing
        edit_option = input("\nDo you want to edit the generated document further? (y/n): ").strip().lower()
        if edit_option == 'y':
            file_to_edit = input("Enter the path to the generated markdown/text file: ").strip()
            iterative_editing_loop(processor, file_to_edit, document_type)
    except Exception as e:
        print(f"An error occurred: {e}")

if __name__ == "__main__":
    main()
